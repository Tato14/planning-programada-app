"""Plantilla del fitxer d'entrada setmanal (buida o amb dades).

Pestanyes: Instruccions, Configuració, Radiòlegs, Demanda, Absències. Inclou
desplegables (activitats, franges, àmbit, codis d'absència, radiòlegs). A Absències
el nom s'escriu (o es tria) a cada fila en lloc d'enllaçar-lo per fórmula: així,
si algú reordena o insereix files a Radiòlegs, les absències no canvien de persona.
"""
from __future__ import annotations

import datetime as dt
from typing import Optional

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .normalize import WEEKDAY_SHORT, FRANGES, AMBITS, fmt_date

DEFAULT_ACTIVITIES = ['RM Neuro', 'RM Body', 'RM MSK', 'RM Cardio', 'RM Miscel·lània',
                      'TC Neuro', 'TC Body', 'TC MSK', 'TC Cardio', 'TC Miscel·lània']

TITLE = Font(bold=True, size=14, color='1F4E79')
NOTE = Font(italic=True, size=9, color='595959')
HEAD_FILL = PatternFill('solid', fgColor='1F4E79')
HEAD_FONT = Font(bold=True, color='FFFFFF')
GROUP_FILL = PatternFill('solid', fgColor='DDEBF7')
GROUP_FONT = Font(bold=True, color='1F4E79')
INPUT_FILL = PatternFill('solid', fgColor='FFFDE7')
THIN = Side(style='thin', color='BFBFBF')
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical='top')

PROFILE_HEADERS = ['Radiòleg', 'Correu', 'Actiu aquesta setmana', 'Mínim setmanal', 'Màxim setmanal',
                   'Centres on pot informar', 'Dies sense activitat', 'Observacions']
DEMAND_HEADERS = ['Centre', 'Data', 'Franja', 'Agenda / equip', 'Activitat', 'Exploracions', 'Àmbit',
                  'Radiòleg fix', 'Observacions']
HEADER_ROW = 4
FIRST = HEADER_ROW + 1
CFG = "'Configuració'"
RAD = "'Radiòlegs'"


def _title(ws, text: str, note: str, width_cols: int):
    ws['A1'] = text
    ws['A1'].font = TITLE
    ws['A2'] = note
    ws['A2'].font = NOTE
    ws['A2'].alignment = Alignment(wrap_text=True, vertical='top')
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max(2, width_cols))
    ws.row_dimensions[2].height = 42


def _header(ws, headers: list, row: int = HEADER_ROW, start_col: int = 1):
    for j, h in enumerate(headers):
        c = ws.cell(row=row, column=start_col + j, value=h)
        c.font = HEAD_FONT
        c.fill = HEAD_FILL
        c.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
        c.border = BOX
    ws.row_dimensions[row].height = 32


def _dv(ws, formula: str, rng: str, title: str, msg: str, strict: bool = True, kind: str = 'list'):
    dv = DataValidation(type=kind, formula1=formula, allow_blank=True, showErrorMessage=strict,
                        errorTitle=title, error=msg[:220], showInputMessage=True, promptTitle=title, prompt=msg[:250])
    if kind != 'list':
        dv.operator = 'greaterThanOrEqual'
    ws.add_data_validation(dv)
    dv.add(rng)
    return dv


def build_input(path: str, week: Optional[dt.date] = None, holidays: Optional[list] = None,
                params: Optional[dict] = None, radiologists: Optional[list] = None,
                activities: Optional[list] = None, demand: Optional[list] = None,
                absences: Optional[dict] = None, blank_rows: int = 80) -> str:
    """radiologists: [{'Radiòleg', 'Correu', 'Actiu aquesta setmana', 'Mínim setmanal', 'Màxim setmanal',
    'Centres on pot informar', 'Dies sense activitat', 'Observacions', 'competències': {activitat: 'S'|'P'}}]
    demand: [{capçalera de DEMAND_HEADERS: valor}]   absences: {nom: {dia_setmana (0-6): codi}}"""
    activities = list(activities or DEFAULT_ACTIVITIES)
    radiologists = radiologists or []
    demand = demand or []
    absences = absences or {}
    params = params or {}
    wb = openpyxl.Workbook()

    # ------------------------------------------------------------ Instruccions
    ws = wb.active
    ws.title = 'Instruccions'
    ws['A1'] = "Entrada setmanal · assignació d'activitat programada TD"
    ws['A1'].font = TITLE
    lines = [
        ('Una vegada', "Omple la pestanya Radiòlegs: una fila per radiòleg. Només cal tocar-la quan hi ha canvis (alta, baixa, canvi de contracte o de preferències)."),
        ('Cada setmana', "Fes una còpia del fitxer, canvia la setmana i els festius a Configuració, omple Demanda i Absències, i puja'l a l'app."),
        ('Mínim i Màxim', "En exploracions per setmana completa. Mínim = obligació contractual (0 si no en té). Màxim = el que pot fer. Si no vol extra, Màxim = Mínim. Els contractes que només cobreixen guàrdies tenen Màxim 0."),
        ('Activitats', "Una columna per activitat. S = la fa · P = la fa i la prefereix · buit = no la fa. Per afegir una activitat, afegeix una columna al final amb el nom (p. ex. 'RM Mama')."),
        ('Actiu aquesta setmana', "Sí o buit = participa. No = queda fora aquesta setmana sense esborrar la fila. Un valor que no s'entén es tracta com a No i surt un avís."),
        ('Centres', "Centres on pot informar (accés al HIS). Buit o 'Tots' = tots. Si n'hi ha, separa'ls amb comes i escriu-los com surten a la Demanda (n'hi ha prou amb el nom curt: 'Bellvitge')."),
        ('Dies sense activitat', "Per a contractes parcials amb dies fixos (p. ex. 'Dj, Dv'). No redueixen el Màxim: serveixen per saber quins dies hi és (ingressats) i per comptar bé les absències."),
        ('Demanda', "Una fila per agenda (centre, dia i franja) o per volum de la setmana (sense data). L'activitat ha de ser una de les columnes de Radiòlegs."),
        ('Àmbit', "Ingressat o Mixt si l'agenda té pacients ingressats: termini de 24 h, només s'assigna a qui treballa aquell dia."),
        ('Radiòleg fix', "Només si hi ha un pacte. Si en són diversos, separa'ls amb punt i coma (la coma ja és dins de 'Cognoms, Nom'). Si el fix no treballa el dia del bloc, es busca substitut (queda marcat); en caps de setmana i festius el pacte es manté."),
        ('Absències', "Una fila per radiòleg amb alguna absència la setmana: V vacances · B baixa · C congrés/formació · A altra absència · G guàrdia en una altra institució · M mig dia · X guàrdia ja assignada (informatiu). Cada setmana, buida les files de la setmana anterior."),
        ("Com s'assigna", "1) Cobrir la demanda (primer la que té ingressats) · 2) Respectar el radiòleg fix · 3) Omplir l'obligació de cadascú · 4) Repartir l'extra en proporció a la capacitat declarada, fins a l'ocupació objectiu · 5) No partir agendes sense necessitat · 6) Activitat preferida."),
    ]
    for i, (k, v) in enumerate(lines, start=3):
        ws.cell(row=i, column=1, value=k).font = Font(bold=True, color='1F4E79')
        c = ws.cell(row=i, column=2, value=v)
        c.alignment = WRAP
        ws.row_dimensions[i].height = 44
    ws.column_dimensions['A'].width = 22
    ws.column_dimensions['B'].width = 110

    # ----------------------------------------------------------- Configuració
    ws = wb.create_sheet('Configuració')
    _title(ws, 'Configuració de la setmana', "Cada setmana només cal canviar la data del dilluns i els festius. La resta són paràmetres de l'assignació (valors per defecte recomanats).", 3)
    _header(ws, ['Paràmetre', 'Valor', 'Notes'])
    cfg_rows = [
        ('Setmana (dilluns)', week, "Data del dilluns de la setmana que es planifica (dd/mm/aaaa)."),
        ('Festius', ', '.join(fmt_date(d) for d in (holidays or [])), "dd/mm/aaaa separats per comes. Només compten els d'aquesta setmana."),
        ('Dies laborables', params.get('Dies laborables', 'Dl-Dv'), 'Dies en què es treballa programada.'),
        ("Ocupació objectiu de l'extra (%)", params.get('Ocupació objectiu', 80), "Si algú declara 50 d'extra, se n'hi proposen fins a 40 abans de passar del 80%."),
        ('Mida mínima de lot (exploracions)', params.get('Mida mínima de lot', 5), "Si es divideix una agenda, cap part per sota d'aquesta mida."),
        ('Màxim de radiòlegs per bloc', params.get('Màxim de radiòlegs per bloc', 4), ''),
        ('Cost de dividir un bloc (exploracions equivalents)', params.get('Cost de dividir', 20), 'Com més alt, menys es divideixen les agendes.'),
        ('Temps màxim de càlcul (s)', params.get('Temps màxim de càlcul', 20), 'Normalment n\'hi ha prou amb 20.'),
    ]
    for i, (k, v, n) in enumerate(cfg_rows, start=FIRST):
        ws.cell(row=i, column=1, value=k).border = BOX
        c = ws.cell(row=i, column=2, value=v)
        c.fill = INPUT_FILL
        c.border = BOX
        if i == FIRST:
            c.number_format = 'dd/mm/yyyy'
        ws.cell(row=i, column=3, value=n).font = NOTE
    ws.column_dimensions['A'].width = 48
    ws.column_dimensions['B'].width = 26
    ws.column_dimensions['C'].width = 80
    _dv(ws, '0', f"B{FIRST}", 'Setmana', 'Data del dilluns (dd/mm/aaaa).', kind='date')

    # -------------------------------------------------------------- Radiòlegs
    ws = wb.create_sheet('Radiòlegs')
    acts_first = len(PROFILE_HEADERS) + 1
    last_col = acts_first + len(activities) - 1
    _title(ws, 'Radiòlegs i preferències', "Una fila per radiòleg. Mínim i Màxim en exploracions per setmana completa. Activitats: S = la fa, P = la fa i la prefereix, buit = no la fa.", last_col)
    ws.cell(row=3, column=1, value='Perfil').font = GROUP_FONT
    ws.cell(row=3, column=1).fill = GROUP_FILL
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=len(PROFILE_HEADERS))
    ws.cell(row=3, column=acts_first, value='Activitats: S = la fa · P = la fa i la prefereix').font = GROUP_FONT
    ws.cell(row=3, column=acts_first).fill = GROUP_FILL
    ws.merge_cells(start_row=3, start_column=acts_first, end_row=3, end_column=max(acts_first, last_col))
    _header(ws, PROFILE_HEADERS + activities)
    n_rows = max(blank_rows, len(radiologists) + 20)
    for i, r in enumerate(radiologists):
        row = FIRST + i
        for j, h in enumerate(PROFILE_HEADERS, start=1):
            ws.cell(row=row, column=j, value=r.get(h))
        comps = r.get('competències', {})
        for j, a in enumerate(activities):
            v = comps.get(a)
            if v:
                ws.cell(row=row, column=acts_first + j, value=v).alignment = Alignment(horizontal='center')
    widths = [24, 28, 12, 10, 10, 30, 14, 26]
    for j, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(j)].width = w
    for j in range(acts_first, last_col + 1):
        ws.column_dimensions[get_column_letter(j)].width = 11
    end_row = FIRST + n_rows - 1
    _dv(ws, '"Sí,No"', f"C{FIRST}:C{end_row}", 'Actiu', 'Sí o No (buit = Sí).')
    _dv(ws, '0', f"D{FIRST}:E{end_row}", 'Capacitat', 'Exploracions per setmana (número ≥ 0).', kind='decimal')
    _dv(ws, '"S,P"', f"{get_column_letter(acts_first)}{FIRST}:{get_column_letter(last_col + 12)}{end_row}",
        'Competència', 'S = la fa · P = la fa i la prefereix · buit = no la fa.')
    ws.freeze_panes = ws.cell(row=FIRST, column=2)

    # ---------------------------------------------------------------- Demanda
    ws = wb.create_sheet('Demanda')
    _title(ws, 'Demanda de la setmana', "Una fila per agenda (centre, dia i franja) o per volum de la setmana (sense data). Activitat: una de les columnes de Radiòlegs. Radiòleg fix: només si hi ha un pacte.", len(DEMAND_HEADERS))
    _header(ws, DEMAND_HEADERS)
    for i, d in enumerate(demand):
        row = FIRST + i
        for j, h in enumerate(DEMAND_HEADERS, start=1):
            c = ws.cell(row=row, column=j, value=d.get(h))
            if isinstance(d.get(h), dt.date):
                c.number_format = 'dd/mm/yyyy'
    for j, w in enumerate([24, 12, 10, 18, 18, 12, 12, 24, 30], start=1):
        ws.column_dimensions[get_column_letter(j)].width = w
    d_end = FIRST + max(blank_rows * 2, len(demand) + 40) - 1
    for r in range(FIRST, d_end + 1):
        ws.cell(row=r, column=2).number_format = 'dd/mm/yyyy'
    dv = _dv(ws, f"{CFG}!$B${FIRST}", f"B{FIRST}:B{d_end}", 'Data', "Una data de la setmana planificada (cal el dilluns a Configuració), o buit si és volum de tota la setmana.", kind='date')
    dv.operator = 'between'
    dv.formula2 = f"{CFG}!$B${FIRST}+6"
    _dv(ws, '"' + ','.join(FRANGES) + '"', f"C{FIRST}:C{d_end}", 'Franja', 'Matí, Tarda, Nit o Tot el dia.')
    act_rng = f"{RAD}!${get_column_letter(acts_first)}${HEADER_ROW}:${get_column_letter(last_col + 12)}${HEADER_ROW}"
    _dv(ws, act_rng, f"E{FIRST}:E{d_end}", 'Activitat', "Ha de ser una de les columnes d'activitat de la pestanya Radiòlegs.")
    dvn = _dv(ws, '0', f"F{FIRST}:F{d_end}", 'Exploracions', 'Nombre enter d\'exploracions (> 0).', kind='whole')
    dvn.operator = 'greaterThan'
    _dv(ws, '"' + ','.join(AMBITS) + '"', f"G{FIRST}:G{d_end}", 'Àmbit', 'Ambulatori (o buit), Ingressat o Mixt.')
    _dv(ws, f"{RAD}!$A${FIRST}:$A${end_row}", f"H{FIRST}:H{d_end}", 'Radiòleg fix',
        "Només si hi ha un pacte. Diversos: separa'ls amb punt i coma.", strict=False)
    ws.freeze_panes = ws.cell(row=FIRST, column=1)

    # -------------------------------------------------------------- Absències
    ws = wb.create_sheet('Absències')
    _title(ws, 'Absències de la setmana', "Només els radiòlegs que tenen alguna absència (desplegable amb els noms de Radiòlegs). V vacances · B baixa · C congrés/formació · A altra absència · G guàrdia en una altra institució · M mig dia · X guàrdia ja assignada (informatiu).", 9)
    for d in range(7):
        c = ws.cell(row=3, column=2 + d, value=f'=IF({CFG}!$B${FIRST}="","",{CFG}!$B${FIRST}+{d})')
        c.number_format = 'dd/mm'
        c.font = NOTE
        c.alignment = Alignment(horizontal='center')
    _header(ws, ['Radiòleg'] + WEEKDAY_SHORT + ['Observacions'])
    row = FIRST
    for name, days in absences.items():
        ws.cell(row=row, column=1, value=name)
        for wd, code in days.items():
            ws.cell(row=row, column=2 + wd, value=code).alignment = Alignment(horizontal='center')
        row += 1
    a_end = FIRST + max(blank_rows, len(absences) + 20) - 1
    ws.column_dimensions['A'].width = 26
    for j in range(2, 9):
        ws.column_dimensions[get_column_letter(j)].width = 7
    ws.column_dimensions['I'].width = 30
    _dv(ws, f"{RAD}!$A${FIRST}:$A${end_row}", f"A{FIRST}:A{a_end}", 'Radiòleg',
        'Tria el radiòleg de la llista (els noms de la pestanya Radiòlegs).', strict=False)
    _dv(ws, '"V,B,C,A,G,M,X"', f"B{FIRST}:H{a_end}", 'Absència',
        'V vacances · B baixa · C congrés/formació · A altra · G guàrdia fora · M mig dia · X guàrdia assignada.')
    ws.freeze_panes = ws.cell(row=FIRST, column=2)

    wb.active = 1
    wb.save(path)
    return path
