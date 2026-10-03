"""Excel de sortida de l'assignació.

Totes les vistes són Taules d'Excel amb nom (tAssignacions, tRadiolegs, tPendents...),
de manera que es poden filtrar a Excel i, més endavant, llegir des de Power Automate
amb l'acció estàndard "List rows present in a table".
"""
from __future__ import annotations

import datetime as dt
from collections import defaultdict
from typing import Optional

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from . import __version__
from .model import PlanResult, ERROR, AVIS, INFO
from .normalize import WEEKDAY_SHORT, week_label

HEADER_FILL = PatternFill('solid', fgColor='1F4E79')
HEADER_FONT = Font(bold=True, color='FFFFFF')
TITLE_FONT = Font(bold=True, size=14, color='1F4E79')
SUB_FONT = Font(bold=True, size=11, color='1F4E79')
WRAP = Alignment(wrap_text=True, vertical='top')
LEVEL_FILL = {ERROR: PatternFill('solid', fgColor='F8CBAD'), AVIS: PatternFill('solid', fgColor='FFF2CC'),
              INFO: PatternFill('solid', fgColor='EDEDED')}
DATE_FMT = 'dd/mm/yyyy'
PCT_FMT = '0%'


def _add_table(ws, name: str, headers: list, rows: list, start_row: int = 1, start_col: int = 1,
               widths: Optional[dict] = None, wrap_cols: tuple = (), pct_cols: tuple = (), freeze: bool = True):
    for j, h in enumerate(headers):
        c = ws.cell(row=start_row, column=start_col + j, value=h)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.alignment = Alignment(wrap_text=True, vertical='center')
    data = rows if rows else [[None] * len(headers)]
    for i, row in enumerate(data, start=1):
        for j, v in enumerate(row):
            c = ws.cell(row=start_row + i, column=start_col + j, value=v)
            if isinstance(v, dt.datetime):
                c.number_format = 'dd/mm/yyyy hh:mm'
            elif isinstance(v, dt.date):
                c.number_format = DATE_FMT
            if headers[j] in pct_cols:
                c.number_format = PCT_FMT
            if headers[j] in wrap_cols:
                c.alignment = WRAP
    end_row = start_row + len(data)
    end_col = start_col + len(headers) - 1
    ref = f"{get_column_letter(start_col)}{start_row}:{get_column_letter(end_col)}{end_row}"
    t = Table(displayName=name, ref=ref)
    t.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
    ws.add_table(t)
    for j, h in enumerate(headers):
        col = get_column_letter(start_col + j)
        w = (widths or {}).get(h)
        if w is None:
            longest = max([len(str(h))] + [len(str(r[j])) for r in rows[:300] if r[j] is not None] or [10])
            w = min(max(9, longest + 2), 45)
        ws.column_dimensions[col].width = w
    if freeze:
        ws.freeze_panes = ws.cell(row=start_row + 1, column=start_col)
    return end_row


def _day(d: Optional[dt.date]) -> str:
    return WEEKDAY_SHORT[d.weekday()] if d else 'Setmana'


def assignment_rows(res: PlanResult) -> tuple:
    bmap = res.blocks_by_id()
    headers = ['Data', 'Dia', 'Franja', 'Centre', 'Agenda / equip', 'Activitat', 'Àmbit', 'Exploracions',
               'Radiòleg', 'Correu', 'Tipus', 'Dins obligació', 'Extra', 'Motiu', 'Avisos', 'Fila demanda', 'ID']
    rows = []
    for a in res.assignments:
        b = bmap[a.block_id]
        r = res.entrada.radiologists[a.radiologist_key]
        rows.append([b.date, _day(b.date), b.franja, b.centre, b.agenda, b.activity_label, b.scope, a.n_exams,
                     r.name, r.email, a.kind, a.oblig_exams, a.extra_exams, a.reason, ' · '.join(a.flags), b.row, a.id])
    return headers, rows


def radiologist_rows(res: PlanResult) -> tuple:
    bmap = res.blocks_by_id()
    acts = res.entrada.activities
    headers = ['Radiòleg', 'Correu', 'Disponibilitat', 'Absències', 'Mínim (setmana)', 'Màxim (setmana)',
               'Assignades', 'Dins obligació', 'Extra', 'Ocupació', "Ocupació de l'extra", 'En activitat preferida',
               'Activitats assignades']
    per = defaultdict(list)
    for a in res.assignments:
        per[a.radiologist_key].append(a)
    rows = []
    for k, rw in sorted(res.radweeks.items(), key=lambda kv: kv[1].radiologist.name):
        r = rw.radiologist
        mine = per.get(k, [])
        n = sum(a.n_exams for a in mine)
        ob = sum(a.oblig_exams for a in mine)
        ex = sum(a.extra_exams for a in mine)
        pref = sum(a.n_exams for a in mine if r.prefers(bmap[a.block_id].activity_key))
        by_act = defaultdict(int)
        for a in mine:
            by_act[bmap[a.block_id].activity_key] += a.n_exams
        acts_txt = ', '.join(f"{acts[ak].label if ak in acts else ak} {v}" for ak, v in sorted(by_act.items(), key=lambda kv: -kv[1]))
        rows.append([r.name, r.email, round(rw.availability, 3), rw.absence_text(), round(rw.min, 1), round(rw.cap, 1),
                     n, ob, ex, round(n / rw.cap, 3) if rw.cap else None,
                     round(max(0, n - rw.min) / rw.extra, 3) if rw.extra > 0 else None,
                     round(pref / n, 3) if n else None, acts_txt])
    return headers, rows


def pending_rows(res: PlanResult) -> tuple:
    bmap = res.blocks_by_id()
    headers = ['Data', 'Dia', 'Franja', 'Centre', 'Agenda / equip', 'Activitat', 'Àmbit', 'Pendents', 'Motius',
               'Candidats', 'Fila demanda']
    rows = []
    for pd_ in res.pendings:
        b = bmap[pd_.block_id]
        rows.append([b.date, _day(b.date), b.franja, b.centre, b.agenda, b.activity_label, b.scope, pd_.n_exams,
                     ' · '.join(pd_.reasons), ', '.join(pd_.candidates), b.row])
    return headers, rows


def summary_by(res: PlanResult, attr: str) -> list:
    """[(valor, demanda, assignades, pendents, radiòlegs)] per activitat o per centre."""
    bmap = res.blocks_by_id()
    dem, asg, rads = defaultdict(int), defaultdict(int), defaultdict(set)
    for b in res.blocks:
        dem[getattr(b, attr)] += b.n_exams
    for a in res.assignments:
        v = getattr(bmap[a.block_id], attr)
        asg[v] += a.n_exams
        rads[v].add(a.radiologist_key)
    out = [(v, dem[v], asg[v], dem[v] - asg[v], len(rads[v])) for v in dem]
    out.sort(key=lambda t: -t[1])
    return out


def write_plan(res: PlanResult, path: str) -> str:
    wb = openpyxl.Workbook()

    # ------------------------------------------------------------------ Resum
    ws = wb.active
    ws.title = 'Resum'
    demand = sum(b.n_exams for b in res.blocks)
    assigned = sum(a.n_exams for a in res.assignments)
    pending = sum(p.n_exams for p in res.pendings)
    ws['A1'] = f"Assignació d'activitat programada · {week_label(res.week)}"
    ws['A1'].font = TITLE_FONT
    ws['A2'] = f"Generat el {res.generated_at:%d/%m/%Y %H:%M} a partir de {res.entrada.source} · versió {__version__}"
    kpis = [('Exploracions demanades', demand), ('Assignades', assigned), ('Pendents', pending),
            ('Cobertura', round(assigned / demand, 3) if demand else None),
            ('Radiòlegs amb activitat', len({a.radiologist_key for a in res.assignments})),
            ('Estat del càlcul', res.stats.status)]
    for i, (k, v) in enumerate(kpis, start=4):
        ws.cell(row=i, column=1, value=k).font = SUB_FONT
        c = ws.cell(row=i, column=2, value=v)
        if k == 'Cobertura':
            c.number_format = PCT_FMT
    errors = [w for w in res.validation if w.level == ERROR]
    ws.cell(row=11, column=1, value='Validació independent').font = SUB_FONT
    ws.cell(row=11, column=2, value='Superada' if not errors else f"{len(errors)} errors: vegeu Avisos")
    r0 = 13
    ws.cell(row=r0, column=1, value='Per activitat').font = SUB_FONT
    rows = [list(t) for t in summary_by(res, 'activity_label')]
    end = _add_table(ws, 'tResumActivitat', ['Activitat', 'Demanda', 'Assignades', 'Pendents', 'Radiòlegs'], rows,
                     start_row=r0 + 1, freeze=False)
    r1 = end + 2
    ws.cell(row=r1, column=1, value='Per centre').font = SUB_FONT
    rows = [list(t) for t in summary_by(res, 'centre')]
    _add_table(ws, 'tResumCentre', ['Centre', 'Demanda', 'Assignades', 'Pendents', 'Radiòlegs'], rows,
               start_row=r1 + 1, freeze=False)
    ws.column_dimensions['A'].width = 34
    ws.column_dimensions['B'].width = 16

    # ------------------------------------------------------------ Assignacions
    ws = wb.create_sheet('Assignacions')
    headers, rows = assignment_rows(res)
    _add_table(ws, 'tAssignacions', headers, rows, widths={'Motiu': 30, 'Avisos': 50, 'Correu': 28, 'Radiòleg': 22,
                                                           'Centre': 22, 'ID': 12},
               wrap_cols=('Avisos', 'Motiu'))

    # ------------------------------------------------------------- Radiòlegs
    ws = wb.create_sheet('Radiòlegs')
    headers, rows = radiologist_rows(res)
    _add_table(ws, 'tRadiolegs', headers, rows, widths={'Activitats assignades': 45, 'Absències': 30, 'Correu': 28},
               wrap_cols=('Activitats assignades', 'Absències'),
               pct_cols=('Disponibilitat', 'Ocupació', "Ocupació de l'extra", 'En activitat preferida'))

    # -------------------------------------------------------------- Pendents
    ws = wb.create_sheet('Pendents')
    headers, rows = pending_rows(res)
    _add_table(ws, 'tPendents', headers, rows, widths={'Motius': 70, 'Candidats': 30}, wrap_cols=('Motius', 'Candidats'))

    # ---------------------------------------------------------------- Avisos
    ws = wb.create_sheet('Avisos')
    order = {ERROR: 0, AVIS: 1, INFO: 2}
    allw = sorted(list(res.validation) + list(res.warnings), key=lambda w: order.get(w.level, 3))
    headers = ['Nivell', 'Categoria', 'Missatge', 'Origen', 'Suggeriment']
    rows = [[w.level, w.category, w.message, w.source, w.suggestion] for w in allw]
    end = _add_table(ws, 'tAvisos', headers, rows, widths={'Missatge': 80, 'Suggeriment': 40, 'Origen': 22},
                     wrap_cols=('Missatge', 'Suggeriment'))
    for i, w in enumerate(allw, start=2):
        ws.cell(row=i, column=1).fill = LEVEL_FILL.get(w.level, LEVEL_FILL[INFO])

    # ------------------------------------------------------------- Paràmetres
    ws = wb.create_sheet('Paràmetres')
    st = res.stats
    rows = [['Setmana planificada', res.week], ['Fitxer d\'entrada', res.entrada.source],
            ['Generat', res.generated_at], ['Versió', __version__]]
    rows += [[k, v] for k, v in res.params.describe()]
    rows += [['Estat del càlcul', st.status], ['Temps de càlcul (s)', round(st.wall_time, 1)],
             ['Marge fins a l\'òptim', round(st.gap, 4) if st.gap is not None else None],
             ['Llavor', st.seed], ['Blocs', st.n_blocks], ['Radiòlegs disponibles', st.n_radiologists]]
    _add_table(ws, 'tParametres', ['Paràmetre', 'Valor'], rows, widths={'Paràmetre': 48, 'Valor': 30})
    for i, (k, _) in enumerate(rows, start=2):
        if k == "Marge fins a l'òptim":
            ws.cell(row=i, column=2).number_format = '0.0%'

    for ws in wb.worksheets:
        ws.page_setup.orientation = 'landscape'
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
    wb.save(path)
    return path
