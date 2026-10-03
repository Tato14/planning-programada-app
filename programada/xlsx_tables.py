"""Lectura tolerant de taules dins de fulls d'Excel.

Les plantilles tenen títol i instruccions a dalt i la capçalera a la fila 3 o 8,
però les persones mouen columnes, n'afegeixen o canvien una mica el nom de la
capçalera. Aquí es localitza la fila de capçalera comparant els textos amb una
llista d'àlies per columna, i es retornen les files com a diccionaris amb claus
canòniques.

També es resol el patró de fórmula de les plantilles (=SI(Radiòlegs!A4="";"";Radiòlegs!A4))
quan el fitxer s'ha generat amb openpyxl i encara no s'ha desat amb Excel (sense
valors en memòria cau).
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from typing import Optional

import openpyxl

from .normalize import norm_key, clean


@dataclass
class Book:
    path: str
    data: openpyxl.Workbook
    formulas: openpyxl.Workbook
    label: str = ''      # nom a mostrar als avisos (el fitxer original, no el temporal)
    unresolved: list = field(default_factory=list)   # [(full, fila, columna, fórmula)] sense valor calculat

    @classmethod
    def open(cls, path) -> 'Book':
        data = openpyxl.load_workbook(path, data_only=True)
        formulas = openpyxl.load_workbook(path, data_only=False)
        return cls(str(path), data, formulas)

    def sheet(self, *names) -> Optional[str]:
        """Retorna el nom real del primer full que coincideixi (sense accents/majúscules)."""
        wanted = [norm_key(n) for n in names]
        for sn in self.data.sheetnames:
            if norm_key(sn) in wanted:
                return sn
        for sn in self.data.sheetnames:
            k = norm_key(sn)
            if any(k.startswith(w) for w in wanted if len(w) >= 5):
                return sn
        return None

    def value(self, sheet: str, row: int, col: int):
        """Valor de la cel·la; si és una fórmula sense valor desat, intenta resoldre-la."""
        v = self.data[sheet].cell(row=row, column=col).value
        if v is not None:
            return v
        f = self.formulas[sheet].cell(row=row, column=col).value
        if isinstance(f, str) and f.startswith('='):
            r = self._resolve_formula(f, depth=0)
            if r is _UNRESOLVED:
                self.unresolved.append((sheet, row, col, f))
                return None
            return r
        return None

    _REF = r"(?:'[^']+'|[A-Za-zÀ-ÿ_][\wÀ-ÿ\.]*)!\$?[A-Z]{1,3}\$?\d+"
    _ONE = re.compile(r"^=\s*(" + _REF + r")\s*$")
    _IF = re.compile(r'^=\s*IF\(\s*(' + _REF + r')\s*=\s*""\s*[,;]\s*""\s*[,;]\s*(' + _REF + r')\s*\)\s*$', re.I)
    _PARTS = re.compile(r"(?:'([^']+)'|([A-Za-zÀ-ÿ_][\wÀ-ÿ\.]*))!\$?([A-Z]{1,3})\$?(\d+)")

    def _resolve_formula(self, f: str, depth: int):
        """Només els dos patrons de les plantilles: =Full!A4 i =IF(Full!A4="","",Full!A4).
        Qualsevol altra fórmula (amb aritmètica, funcions...) no es pot calcular sense Excel."""
        if depth > 3:
            return _UNRESOLVED
        m = self._ONE.match(f)
        ref = m.group(1) if m else None
        if ref is None:
            m = self._IF.match(f)
            if m and m.group(1).replace('$', '') == m.group(2).replace('$', ''):
                ref = m.group(2)
        if ref is None:
            return _UNRESOLVED
        sheet_name, col_letters, row_s = (lambda g: (g[0] or g[1], g[2], g[3]))(self._PARTS.match(ref).groups())
        real = None
        for sn in self.data.sheetnames:
            if sn == sheet_name or norm_key(sn) == norm_key(sheet_name):
                real = sn
                break
        if real is None:
            return _UNRESOLVED
        from openpyxl.utils import column_index_from_string
        col = column_index_from_string(col_letters)
        row = int(row_s)
        v = self.data[real].cell(row=row, column=col).value
        if v is None:
            f2 = self.formulas[real].cell(row=row, column=col).value
            if isinstance(f2, str) and f2.startswith('='):
                return self._resolve_formula(f2, depth + 1)
            return None
        return v if v != '' else None


_UNRESOLVED = object()


def _match_score(header_key: str, aliases: list[str]) -> int:
    best = 0
    for a in aliases:
        ak = norm_key(a)
        if not ak:
            continue
        if header_key == ak:
            best = max(best, 100 + len(ak))
        elif len(ak) >= 4 and header_key.startswith(ak):
            best = max(best, len(ak))
    return best


def find_header(book: Book, sheet: str, spec: dict, max_scan: int = 40,
                required: Optional[list] = None, accept=None) -> tuple[Optional[int], dict]:
    """Localitza la fila de capçalera.

    spec: {clau_canònica: [àlies, ...]}. Entre les files candidates guanya la que té més
    coincidències exactes, després més coincidències i després més puntuació, de manera
    que un títol que comença com una capçalera ("Radiòlegs absents...") no la desplaça.
    accept(fila) -> bool permet exigir alguna cosa més a la fila (p. ex. columnes de dies).
    Retorna (fila, {clau_canònica: columna}).
    """
    ws = book.data[sheet]
    best_row, best_map, best_key = None, {}, None
    max_col = min(ws.max_column, 80)
    for r in range(1, min(ws.max_row, max_scan) + 1):
        colmap = {}
        scores = {}
        for c in range(1, max_col + 1):
            v = ws.cell(row=r, column=c).value
            if v is None or not isinstance(v, str):
                continue
            hk = norm_key(v)
            if not hk:
                continue
            for canon, aliases in spec.items():
                sc = _match_score(hk, aliases)
                if sc and sc > scores.get(canon, 0):
                    scores[canon] = sc
                    colmap[canon] = c
        # Una mateixa columna no pot servir per a dues claus: es queda la de millor puntuació
        used = {}
        for canon, c in sorted(colmap.items(), key=lambda kv: -scores[kv[0]]):
            if c not in used:
                used[c] = canon
        colmap = {canon: c for c, canon in used.items()}
        if not colmap:
            continue
        if required and not all(k in colmap for k in required):
            continue
        if accept is not None and not accept(r):
            continue
        key = (sum(1 for k in colmap if scores[k] >= 100), len(colmap), sum(scores[k] for k in colmap))
        if best_key is None or key > best_key:
            best_row, best_map, best_key = r, colmap, key
    return best_row, best_map


def read_table(book: Book, sheet: str, spec: dict, required: Optional[list] = None,
               key_cols: Optional[list] = None, max_scan: int = 40) -> tuple[list[dict], Optional[int], dict]:
    """Llegeix les files de la taula sota la capçalera.

    key_cols: si s'indica, una fila es considera buida (i s'ignora) quan totes
    aquestes columnes són buides.
    Retorna (files, fila_capçalera, mapa_columnes). Cada fila inclou '_row'.
    """
    hrow, colmap = find_header(book, sheet, spec, max_scan=max_scan, required=required)
    if hrow is None:
        return [], None, {}
    ws = book.data[sheet]
    rows = []
    key_cols = key_cols or list(colmap.keys())
    empty_streak = 0
    for r in range(hrow + 1, ws.max_row + 1):
        rec = {'_row': r}
        for canon, c in colmap.items():
            rec[canon] = book.value(sheet, r, c)
        if all(_is_empty(rec.get(k)) for k in key_cols if k in colmap):
            empty_streak += 1
            if empty_streak > 200:
                break
            continue
        empty_streak = 0
        rows.append(rec)
    return rows, hrow, colmap


def _is_empty(v) -> bool:
    return v is None or (isinstance(v, str) and v.strip() == '')


def extra_columns(book: Book, sheet: str, header_row: int, known_cols: set) -> dict:
    """Columnes de la capçalera no reconegudes: {columna: text_capçalera}."""
    ws = book.data[sheet]
    out = {}
    for c in range(1, min(ws.max_column, 120) + 1):
        if c in known_cols:
            continue
        v = ws.cell(row=header_row, column=c).value
        if v is not None and clean(v):
            out[c] = clean(v)
    return out


def label_value(book: Book, sheet: str, labels: list[str], max_rows: int = 12, max_cols: int = 12):
    """Busca una etiqueta tipus 'Centre:' i retorna el valor de la cel·la de la dreta."""
    ws = book.data[sheet]
    keys = [norm_key(l) for l in labels]
    for r in range(1, min(ws.max_row, max_rows) + 1):
        for c in range(1, min(ws.max_column, max_cols) + 1):
            v = ws.cell(row=r, column=c).value
            if isinstance(v, str) and norm_key(v) in keys:
                for cc in range(c + 1, min(ws.max_column, max_cols + 4) + 1):
                    vv = book.value(sheet, r, cc)
                    if vv is not None and not (isinstance(vv, str) and vv.strip() == ''):
                        return vv
                return None
    return None
