"""Lectura del fitxer d'entrada setmanal (una sola plantilla amb quatre pestanyes).

  Configuració  setmana, festius i paràmetres (tot opcional; hi ha valors per defecte)
  Radiòlegs     un radiòleg per fila: capacitat, centres, dies sense activitat i una
                columna per activitat amb S (la fa) o P (la fa i la prefereix)
  Demanda       un bloc per fila: centre, data, franja, agenda, activitat,
                exploracions, àmbit i, si n'hi ha, el radiòleg fix pactat
  Absències     graella de la setmana amb codis V, B, C, A, G, M, X (opcional)

La lectura és tolerant: busca la fila de capçalera per àlies, accepta columnes
mogudes i textos amb o sense accents, i deixa constància de tot el que no entén a
la llista d'avisos (amb pestanya i fila) en lloc d'aturar-se.
"""
from __future__ import annotations

import datetime as dt
import re
from collections import Counter, defaultdict
from pathlib import Path
from typing import Optional

from .model import (Entrada, Params, Activity, Radiologist, Block, DataWarning, ERROR, AVIS, INFO,
                    centre_matches)
from .normalize import (norm_key, clean, fold, words, split_list, parse_active, parse_number, parse_date,
                        parse_date_list, parse_weekday, parse_weekday_list, parse_franja, parse_ambit,
                        monday_of, fmt_date, AMB, weekday_list_text, round_half_up)
from openpyxl.utils import get_column_letter

from .xlsx_tables import Book, read_table, find_header
from .availability import parse_absence_code, CODES_HELP

SPEC_RADIOLEGS = {
    'name': ['Radiòleg', 'Nom (Cognoms, Nom)', 'Nom', 'Professional', 'Cognoms, Nom', 'Radiólogo'],
    'email': ['Correu', 'Correu electrònic', 'Email', 'E-mail', 'Correo'],
    'active': ['Actiu aquesta setmana', 'Actiu', 'Activo'],
    'min': ['Mínim setmanal', 'Mínim setmanal (obligació)', 'Mínim', 'Obligació setmanal', 'Obligació'],
    'max': ['Màxim setmanal', 'Màxim setmanal (capacitat)', 'Màxim', 'Capacitat setmanal', 'Capacitat'],
    'centres': ['Centres on pot informar', 'Centres', 'Accés HIS', 'Accés a centres'],
    'days_off': ['Dies sense activitat', 'Dies que no treballa', 'Dies lliures'],
    'notes': ['Observacions', 'Notes', 'Comentaris'],
}
SPEC_DEMANDA = {
    'centre': ['Centre', 'Hospital', 'Centro'],
    'date': ['Data', 'Fecha', 'Dia'],
    'franja': ['Franja', 'Torn', 'Horari'],
    'agenda': ['Agenda / equip', 'Agenda', 'Equip / sala', 'Equip', 'Sala'],
    'activity': ['Activitat', "Tipus d'activitat", 'Actividad'],
    'n': ['Exploracions', 'Nº exploracions', 'N exploracions', "Nombre d'exploracions", 'Proves', 'Volum'],
    'scope': ['Àmbit', 'Tipus de pacient', 'Ámbito'],
    'fixed': ['Radiòleg fix', 'Radiòleg fix pactat', 'Radiòleg assignat', 'Fix'],
    'notes': ['Observacions', 'Notes', 'Comentaris'],
}
CONFIG_KEYS = {
    'week': ['Setmana (dilluns)', 'Setmana', 'Dilluns de la setmana', 'Semana'],
    'holidays': ['Festius', 'Festius de la setmana', 'Festivos'],
    'working_days': ['Dies laborables', 'Dies de treball'],
    'target': ["Ocupació objectiu de l'extra", 'Ocupació objectiu', 'Objectiu'],
    'min_chunk': ['Mida mínima de lot', 'Lot mínim'],
    'max_rads': ['Màxim de radiòlegs per bloc', 'Radiòlegs per bloc'],
    'split_cost': ['Cost de dividir un bloc', 'Cost de dividir'],
    'time_limit': ['Temps màxim de càlcul', 'Temps de càlcul'],
}
_COMP_YES = {'s', 'si', 'x', '1', 'y', 'yes', 'ok', 'fa'}
_COMP_PREF = {'p', 'pref', 'preferent', 'preferida', 'preferit', 'preferencia'}
_COMP_NO = {'n', 'no', '0', '-'}


_CHECKS = ('✓', '✔', '☑', '✅')


def parse_competence(v) -> Optional[str]:
    """'S', 'P', '' (no la fa) o None (valor no interpretable).

    Accepta anotacions després del codi ('S (només matins)', 'S*', 'P - preferent') i
    marques de verificació (✓). La resta del text no canvia el codi."""
    if v is None:
        return ''
    if isinstance(v, bool):
        return 'S' if v else ''
    if isinstance(v, (int, float)):
        return 'S' if v == 1 else ('' if v == 0 else None)
    raw = str(v).strip()
    if raw == '' or raw in ('-', '–', '—', '·', 'N/A', 'n/a'):
        return ''
    if raw[0] in _CHECKS:
        return 'S'
    if re.fullmatch(r'\d+([.,]\d+)?', raw):           # un número (p. ex. una dedicació '0,8') no és un codi
        return {'1': 'S', '0': ''}.get(raw)
    head = re.split(r'[(*,;/\-]', raw, maxsplit=1)[0]
    f = norm_key(head) or norm_key(raw)
    if f in _COMP_YES:
        return 'S'
    if f in _COMP_PREF or f.startswith('prefer'):
        return 'P'
    if f in _COMP_NO:
        return ''
    return None


def _numeric(v) -> bool:
    return (isinstance(v, (int, float)) and not isinstance(v, bool)) or bool(re.fullmatch(r'\s*\d+([.,]\d+)?\s*', str(v)))


_ACTIVITY_HEADER = re.compile(r'^(rm|rmn|tc|tac|rx|eco|us|mam|mamo|pet|dxa|mn|angio|intervenc)\b', re.I)
_ALL_CENTRES = {'tots', 'totes', 'totselscentres', 'totselscentre', 'all', 'qualsevol', 'tot'}


def _src(sheet: str, row: int) -> str:
    return f"{sheet}, fila {row}"


class NameIndex:
    """Troba un radiòleg pel nom tal com l'escriu una persona: 'Alpha, Anna', 'Anna Alpha' o 'Alpha'."""

    def __init__(self, radiologists: dict):
        self.rads = list(radiologists.values())
        self.by_key = {r.key: r for r in self.rads}

    def find(self, text) -> tuple:
        """(radiòleg o None, nota)."""
        k = norm_key(text)
        if not k:
            return None, ''
        if k in self.by_key:
            return self.by_key[k], ''
        tw = words(text)
        same = [r for r in self.rads if words(r.name) == tw]
        if len(same) == 1:
            return same[0], ''
        part = [r for r in self.rads if tw and tw <= words(r.name)]
        if len(part) == 1:
            return part[0], f"'{clean(text)}' s'ha interpretat com {part[0].name}"
        if len(part) > 1:
            return None, f"'{clean(text)}' pot ser {', '.join(r.name for r in part[:4])}: escriu el nom complet"
        return None, f"'{clean(text)}' no és a la pestanya Radiòlegs"


# ---------------------------------------------------------------------------
# Configuració
# ---------------------------------------------------------------------------

def _match_config_key(label) -> Optional[str]:
    k = norm_key(label)
    if not k:
        return None
    best, best_len = None, 0
    for canon, aliases in CONFIG_KEYS.items():
        for a in aliases:
            ak = norm_key(a)
            if (k == ak or k.startswith(ak)) and len(ak) > best_len:
                best, best_len = canon, len(ak)
    return best


def _read_config(book: Book, warnings: list) -> tuple:
    p = Params()
    week = None
    sheet = book.sheet('Configuració', 'Configuracio', 'Paràmetres')
    if not sheet:
        warnings.append(DataWarning(INFO, 'Configuració', "No hi ha pestanya Configuració: s'usen els paràmetres per defecte i la setmana es dedueix de les dates de la demanda."))
        return p, week
    ws = book.data[sheet]
    seen = set()
    for r in range(1, min(ws.max_row, 60) + 1):
        for c in range(1, min(ws.max_column, 3) + 1):
            canon = _match_config_key(ws.cell(row=r, column=c).value) if isinstance(ws.cell(row=r, column=c).value, str) else None
            if not canon or canon in seen:
                continue
            v = book.value(sheet, r, c + 1)
            seen.add(canon)
            src = _src(sheet, r)
            if v is None or (isinstance(v, str) and not v.strip()):
                break
            if canon == 'week':
                d = parse_date(v)
                if d is None:
                    warnings.append(DataWarning(AVIS, 'Configuració', f"Setmana '{clean(v)}' no és una data.", src, 'Escriu la data del dilluns (dd/mm/aaaa).'))
                else:
                    if d.weekday() != 0:
                        warnings.append(DataWarning(INFO, 'Configuració', f"{fmt_date(d)} no és dilluns: es planifica la setmana del {fmt_date(monday_of(d))}.", src))
                    week = monday_of(d)
            elif canon == 'holidays':
                p.holidays = parse_date_list(v)
                if not p.holidays:
                    warnings.append(DataWarning(AVIS, 'Configuració', f"Festius '{clean(v)}' no s'han pogut llegir.", src, 'Format dd/mm/aaaa separat per comes.'))
            elif canon == 'working_days':
                wd = parse_weekday_list(v)
                if wd:
                    p.working_days = wd
                else:
                    warnings.append(DataWarning(AVIS, 'Configuració', f"Dies laborables '{clean(v)}' no s'han pogut llegir: s'usa Dl-Dv.", src))
            else:
                num = parse_number(v)
                if num is None:
                    warnings.append(DataWarning(AVIS, 'Configuració', f"Valor '{clean(v)}' no numèric per a '{clean(ws.cell(row=r, column=c).value)}': s'usa el valor per defecte.", src))
                else:
                    attr = {'target': 'target_utilisation', 'min_chunk': 'min_chunk', 'max_rads': 'max_rads_per_block',
                            'split_cost': 'split_cost_exams', 'time_limit': 'time_limit_s'}[canon]
                    val = num / 100 if canon == 'target' and num > 1 else num
                    lo, hi = Params.LIMITS[attr]
                    if not (lo <= val <= hi):
                        shown = (lambda x: f"{x * 100:g}%") if canon == 'target' else (lambda x: f"{x:g}")
                        warnings.append(DataWarning(AVIS, 'Configuració', f"'{clean(ws.cell(row=r, column=c).value)}' = {clean(v)} fora del rang ({shown(lo)}-{shown(hi)}): s'ajusta al límit.", src))
                    setattr(p, attr, val)
            break
    return p.clamp(), week


# ---------------------------------------------------------------------------
# Radiòlegs i activitats
# ---------------------------------------------------------------------------

def _read_radiologists(book: Book, warnings: list) -> tuple:
    sheet = book.sheet('Radiòlegs', 'Radiolegs', 'Professionals', 'Radiòleg')
    if not sheet:
        warnings.append(DataWarning(ERROR, 'Radiòlegs', "Falta la pestanya Radiòlegs.", book.label, 'Fes servir la plantilla d\'entrada.'))
        return {}, {}
    rows, hrow, cmap = read_table(book, sheet, SPEC_RADIOLEGS, required=['name'], key_cols=['name'])
    if hrow is None:
        warnings.append(DataWarning(ERROR, 'Radiòlegs', "No es troba la capçalera (cal una columna 'Radiòleg').", sheet))
        return {}, {}
    ws = book.data[sheet]
    known_cols = set(cmap.values())
    # Columnes de competències. Senyal fort: la fila de grup de la plantilla ("Activitats: ...")
    # damunt de la capçalera marca on comencen. Sense aquesta fila: capçalera amb aspecte
    # d'activitat (RM, TC, RX, Eco...) o valors majoritàriament S / P.
    group_start = group_end = None
    if hrow > 1:
        for c in range(1, min(ws.max_column, 120) + 1):
            g = ws.cell(row=hrow - 1, column=c).value
            if isinstance(g, str) and norm_key(g).startswith('activitat'):
                group_start = group_end = c
                for mr in ws.merged_cells.ranges:
                    if mr.min_row <= hrow - 1 <= mr.max_row and mr.min_col <= c <= mr.max_col:
                        group_end = mr.max_col
                break
    activities, act_cols = {}, {}
    for c in range(1, min(ws.max_column, 120) + 1):
        if c in known_cols:
            continue
        h = ws.cell(row=hrow, column=c).value
        if h is None or not clean(h):
            continue
        label = clean(h)
        vals = [book.value(sheet, rec['_row'], c) for rec in rows]
        filled = [v for v in vals if v is not None and clean(v)]
        # Els números (1, 0,8...) no són prova d'activitat: una columna "Dedicació" no ho és
        ok = [v for v in filled if not _numeric(v) and parse_competence(v) is not None]
        if group_start is not None and group_start <= c <= group_end:
            is_act = True                       # sota l'etiqueta "Activitats" de la plantilla
        elif group_start is not None and c < group_start:
            is_act = False                      # a l'esquerra de les activitats: és perfil
        elif _ACTIVITY_HEADER.match(fold(label)):
            is_act = True
        else:
            is_act = not filled or len(ok) >= 0.7 * len(filled)
        if not is_act:
            example = next((v for v in filled if parse_competence(v) is None), filled[0] if filled else '')
            warnings.append(DataWarning(INFO, 'Radiòlegs', f"Columna '{label}' ignorada: no sembla una activitat (valors com '{clean(example)}' en lloc de S / P).", _src(sheet, hrow),
                                        "Si és una activitat, posa-hi S o P i el nom amb la modalitat (p. ex. 'RM Mama')."))
            continue
        key = norm_key(label)
        if key in activities:
            warnings.append(DataWarning(AVIS, 'Radiòlegs', f"Activitat '{label}' repetida: es fa servir la primera columna.", _src(sheet, hrow)))
            continue
        activities[key] = Activity(key=key, label=label)
        act_cols[key] = c
    if not activities:
        warnings.append(DataWarning(ERROR, 'Radiòlegs', "No hi ha cap columna d'activitat (RM Neuro, TC Body...) amb S / P.", _src(sheet, hrow)))

    rads = {}
    for rec in rows:
        row = rec['_row']
        name = clean(rec.get('name'))
        if not name:
            continue
        key = norm_key(name)
        src = _src(sheet, row)
        if key in rads:
            warnings.append(DataWarning(AVIS, 'Radiòlegs', f"{name} apareix dues vegades: es fa servir la fila {rads[key].row} i s'ignora la fila {row}.", src,
                                        'Esborra la fila que sobra.'))
            continue
        mx = parse_number(rec.get('max'))
        mn = parse_number(rec.get('min'))
        if rec.get('max') not in (None, '') and mx is None:
            warnings.append(DataWarning(AVIS, 'Radiòlegs', f"{name}: Màxim '{clean(rec.get('max'))}' no és un número.", src))
        if rec.get('min') not in (None, '') and mn is None:
            warnings.append(DataWarning(AVIS, 'Radiòlegs', f"{name}: Mínim '{clean(rec.get('min'))}' no és un número.", src))
        mx = max(0.0, mx or 0.0)
        mn = max(0.0, mn or 0.0)
        if mn > mx:
            warnings.append(DataWarning(AVIS, 'Radiòlegs', f"{name}: el Mínim ({mn:g}) és més gran que el Màxim ({mx:g}); es pren el Màxim com a mínim.", src))
            mn = mx
        days_off = parse_weekday_list(rec.get('days_off'))
        if days_off is None:
            warnings.append(DataWarning(AVIS, 'Radiòlegs', f"{name}: dies sense activitat '{clean(rec.get('days_off'))}' no s'entenen.", src, 'Format: Dv, o Dl, Dc'))
            days_off = []
        comps = {}
        for akey, c in act_cols.items():
            v = book.value(sheet, row, c)
            comp = parse_competence(v)
            if comp is None:
                warnings.append(DataWarning(AVIS, 'Radiòlegs', f"{name}: valor '{clean(v)}' a {activities[akey].label} no és S ni P; es considera que no la fa.", src))
            elif comp:
                comps[akey] = comp
        active = parse_active(rec.get('active'))
        if active is None:
            warnings.append(DataWarning(AVIS, 'Radiòlegs', f"{name}: 'Actiu' = '{clean(rec.get('active'))}' no s'entén; es considera inactiu aquesta setmana.", src, 'Sí o No.'))
            active = False
        centres = split_list(rec.get('centres'))
        if any(norm_key(t) in _ALL_CENTRES for t in centres):
            centres = []
        r = Radiologist(name=name, key=key, email=clean(rec.get('email')),
                        active=active, weekly_min=mn, weekly_max=mx,
                        centres=centres, days_off=days_off,
                        notes=clean(rec.get('notes')), competences=comps, row=row)
        if r.active and mx <= 0:
            warnings.append(DataWarning(INFO, 'Radiòlegs', f"{name}: Màxim 0, no rebrà activitat programada.", src))
        if r.active and not comps:
            warnings.append(DataWarning(AVIS, 'Radiòlegs', f"{name}: no té cap activitat marcada amb S o P.", src))
        rads[key] = r
    if not rads:
        warnings.append(DataWarning(ERROR, 'Radiòlegs', 'La pestanya Radiòlegs no té cap radiòleg.', sheet))
    return activities, rads


# ---------------------------------------------------------------------------
# Demanda
# ---------------------------------------------------------------------------

def _resolve_activity(text, activities: dict) -> Optional[Activity]:
    k = norm_key(text)
    if k in activities:
        return activities[k]
    tw = words(text)
    same = [a for a in activities.values() if words(a.label) == tw]
    return same[0] if len(same) == 1 else None


def _read_demand(book: Book, activities: dict, names: NameIndex, week: Optional[dt.date], warnings: list) -> tuple:
    sheet = book.sheet('Demanda', 'Demanda setmanal', 'Blocs')
    if not sheet:
        warnings.append(DataWarning(ERROR, 'Demanda', 'Falta la pestanya Demanda.', book.label))
        return [], week
    rows, hrow, cmap = read_table(book, sheet, SPEC_DEMANDA, required=['activity', 'n'],
                                  key_cols=['centre', 'activity', 'n', 'date'])
    if hrow is None:
        warnings.append(DataWarning(ERROR, 'Demanda', "No es troba la capçalera (calen les columnes 'Activitat' i 'Exploracions').", sheet))
        return [], week
    if not rows:
        warnings.append(DataWarning(ERROR, 'Demanda', 'La pestanya Demanda és buida.', sheet))
        return [], week
    parsed = []
    for rec in rows:
        raw_date = rec.get('date')
        d = parse_date(raw_date)
        wd = None
        if d is None and isinstance(raw_date, str) and raw_date.strip():
            # Dies de la setmana esmentats al text ('Dimecres', 'Dimecres 21/10', 'Dilluns i dimarts (tarda)')
            days = sorted({x for x in (parse_weekday(t) for t in re.split(r'[^a-z]+', fold(raw_date)) if len(t) >= 2) if x is not None})
            if len(days) == 1:
                wd = days[0]
            elif days:
                wd = 'diversos'

        parsed.append((rec, d, wd))
    if week is None:
        dates = [d for _, d, _ in parsed if d]
        if not dates:
            warnings.append(DataWarning(ERROR, 'Configuració', 'No hi ha setmana a Configuració ni dates a la Demanda.', sheet,
                                        'Escriu el dilluns de la setmana a la pestanya Configuració.'))
            return [], None
        week = Counter(monday_of(x) for x in dates).most_common(1)[0][0]
        warnings.append(DataWarning(AVIS, 'Configuració', f"Sense setmana a Configuració: es dedueix de les dates de la demanda ({fmt_date(week)}).", sheet))
    end = week + dt.timedelta(days=6)
    blocks = []
    for rec, d, wd in parsed:
        row = rec['_row']
        src = _src(sheet, row)
        centre = clean(rec.get('centre'))
        act_text = clean(rec.get('activity'))
        n_raw = rec.get('n')
        n = parse_number(n_raw)
        if n is not None and abs(n - round_half_up(n)) > 1e-9:
            warnings.append(DataWarning(INFO, 'Demanda', f"{n:g} exploracions arrodonides a {round_half_up(n)}.", src))
            n = round_half_up(n)
        if n is None or n <= 0:
            warnings.append(DataWarning(AVIS, 'Demanda', f"Fila sense nombre d'exploracions vàlid ('{clean(n_raw)}'): ignorada.", src))
            continue
        n = int(n)
        if not centre:
            warnings.append(DataWarning(AVIS, 'Demanda', 'Fila sense centre: ignorada.', src))
            continue
        if wd == 'diversos':
            warnings.append(DataWarning(AVIS, 'Demanda', f"Data '{clean(rec.get('date'))}' amb diversos dies: es tracta com a volum de la setmana.", src,
                                        'Fes una fila per dia si cal que vagi a qui treballa cada dia.'))
        elif wd is not None:
            d = week + dt.timedelta(days=wd)
        elif d is None and rec.get('date') not in (None, ''):
            warnings.append(DataWarning(AVIS, 'Demanda', f"Data '{clean(rec.get('date'))}' no s'entén: es tracta com a volum de la setmana.", src))
        if d is not None and not (week <= d <= end):
            warnings.append(DataWarning(AVIS, 'Demanda', f"Data {fmt_date(d)} fora de la setmana planificada ({fmt_date(week)}-{fmt_date(end)}): fila ignorada.", src))
            continue
        act = _resolve_activity(act_text, activities) if act_text else None
        if act is None:
            warnings.append(DataWarning(AVIS, 'Demanda', f"Activitat '{act_text or '(buida)'}' no és cap columna de la pestanya Radiòlegs: quedarà pendent.", src,
                                        'Corregeix el nom o afegeix la columna a Radiòlegs.'))
        scope = parse_ambit(rec.get('scope'))
        if scope is None:
            warnings.append(DataWarning(AVIS, 'Demanda', f"Àmbit '{clean(rec.get('scope'))}' no s'entén: es tracta com a ambulatori.", src, 'Ambulatori, Ingressat o Mixt.'))
            scope = AMB
        franja_raw = clean(rec.get('franja'))
        franja = parse_franja(franja_raw) or franja_raw
        fixed_keys, fixed_names = [], []
        for part in _split_names(rec.get('fixed')):
            r, note = names.find(part)
            if r is None:
                warnings.append(DataWarning(AVIS, 'Demanda', f"Radiòleg fix {note}: el bloc s'assigna sense fix.", src))
                continue
            if note:
                warnings.append(DataWarning(INFO, 'Demanda', f"Radiòleg fix {note}.", src))
            if r.key not in fixed_keys:
                fixed_keys.append(r.key)
                fixed_names.append(r.name)
        blocks.append(Block(id=f"F{row:03d}", row=row, week=week, date=d, franja=franja, centre=centre,
                            agenda=clean(rec.get('agenda')), activity_key=act.key if act else '',
                            activity_label=act.label if act else act_text, n_exams=n, scope=scope or AMB,
                            fixed_keys=fixed_keys, fixed_names=fixed_names, notes=clean(rec.get('notes'))))
    if not blocks:
        warnings.append(DataWarning(ERROR, 'Demanda', 'La pestanya Demanda no té cap fila vàlida.', sheet))
    # Possibles duplicats (mateix centre, dia, franja, agenda i activitat)
    seen = defaultdict(list)
    for b in blocks:
        seen[(norm_key(b.centre), b.date, b.franja, norm_key(b.agenda), b.activity_key)].append(b.row)
    for k, rws in seen.items():
        if len(rws) > 1:
            warnings.append(DataWarning(INFO, 'Demanda', f"Files {', '.join(map(str, rws))} tenen el mateix centre, dia, franja, agenda i activitat: comprova que no sigui un duplicat.", sheet))
    return blocks, week


def _split_names(v) -> list:
    """Diversos radiòlegs fixos se separen amb ';', '/' o salt de línia (la coma és dins de 'Cognoms, Nom')."""
    if v is None:
        return []
    return [clean(p) for p in re.split(r'[;/\n]+', str(v)) if clean(p)]


# ---------------------------------------------------------------------------
# Absències
# ---------------------------------------------------------------------------

def _read_absences(book: Book, names: NameIndex, week: dt.date, warnings: list) -> dict:
    sheet = book.sheet('Absències', 'Absencies', 'Vacances', 'Disponibilitat')
    if not sheet:
        warnings.append(DataWarning(INFO, 'Absències', "No hi ha pestanya Absències: tothom es considera disponible tota la setmana."))
        return {}
    ws = book.data[sheet]

    def has_day_columns(r):
        n = 0
        for c in range(1, min(ws.max_column, 40) + 1):
            h = ws.cell(row=r, column=c).value
            if parse_date(h) is not None or (isinstance(h, str) and h.strip() and parse_weekday(h.split()[0]) is not None):
                n += 1
        return n >= 3

    hrow, cmap = find_header(book, sheet, {'name': SPEC_RADIOLEGS['name'], 'notes': SPEC_RADIOLEGS['notes']},
                             required=['name'], accept=has_day_columns)
    if hrow is None:
        warnings.append(DataWarning(AVIS, 'Absències', "No es troba la capçalera (cal una columna 'Radiòleg'): absències ignorades.", sheet))
        return {}
    day_cols = {}
    for c in range(1, min(ws.max_column, 40) + 1):
        if c in cmap.values():
            continue
        h = ws.cell(row=hrow, column=c).value
        d = parse_date(h)
        if d is not None:
            if not (week <= d <= week + dt.timedelta(days=6)):
                warnings.append(DataWarning(AVIS, 'Absències', f"La columna del {fmt_date(d)} no és de la setmana planificada: ignorada.", _src(sheet, hrow)))
                continue
            day_cols[c] = d
            continue
        if isinstance(h, str) and h.strip():
            wd = parse_weekday(h.split()[0])
            if wd is not None:
                day_cols[c] = week + dt.timedelta(days=wd)
    if not day_cols:
        warnings.append(DataWarning(AVIS, 'Absències', 'No es troben les columnes dels dies (Dl, Dt, Dc...): absències ignorades.', _src(sheet, hrow)))
        return {}
    out = defaultdict(dict)
    for r in range(hrow + 1, ws.max_row + 1):
        name = book.value(sheet, r, cmap['name'])
        if name is None or not clean(name):
            continue
        cells = {c: book.value(sheet, r, c) for c in day_cols}
        if not any(v is not None and clean(v) for v in cells.values()):
            continue
        rad, note = names.find(name)
        src = _src(sheet, r)
        if rad is None:
            warnings.append(DataWarning(AVIS, 'Absències', f"Radiòleg {note}: absències ignorades.", src))
            continue
        if note:
            warnings.append(DataWarning(INFO, 'Absències', note, src))
        for c, v in cells.items():
            code, ok = parse_absence_code(v)
            if not code:
                continue
            if not ok:
                warnings.append(DataWarning(AVIS, 'Absències', f"{rad.name}: codi '{clean(v)}' no reconegut el {fmt_date(day_cols[c])}; es tracta com a absència.", src, CODES_HELP))
            out[rad.key][day_cols[c]] = code
    return dict(out)


# ---------------------------------------------------------------------------
# Comprovacions creuades
# ---------------------------------------------------------------------------

def _cross_checks(rads: dict, activities: dict, blocks: list, week: dt.date, p: Params, warnings: list):
    centres = sorted({b.centre for b in blocks})
    for r in rads.values():
        if not r.active:
            continue
        for t in r.centres:
            hits = [c for c in centres if centre_matches(t, c)]
            if len(hits) > 1:
                warnings.append(DataWarning(AVIS, 'Radiòlegs', f"{r.name}: el centre '{t}' coincideix amb {len(hits)} centres de la demanda ({', '.join(hits[:4])}).",
                                            f"Radiòlegs, fila {r.row}", 'Escriu el nom del centre tal com surt a la Demanda.'))
            elif not hits and centres:
                warnings.append(DataWarning(INFO, 'Radiòlegs', f"{r.name}: el centre '{t}' no surt a la demanda d'aquesta setmana.", f"Radiòlegs, fila {r.row}"))
    by_act = defaultdict(int)
    for b in blocks:
        if b.activity_key:
            by_act[b.activity_key] += b.n_exams
    for akey, n in sorted(by_act.items()):
        if not any(r.active and r.can_do(akey) for r in rads.values()):
            warnings.append(DataWarning(AVIS, 'Demanda', f"{activities[akey].label}: {n} exploracions i cap radiòleg actiu que la faci.", 'Demanda'))
    for b in blocks:
        for k in b.fixed_keys:
            if not rads[k].active:
                warnings.append(DataWarning(INFO, 'Demanda', f"{rads[k].name} és el fix de {b.label()} però no és actiu aquesta setmana: es busca substitut.", f"Demanda, fila {b.row}"))
    outside = [d for d in p.holidays if not (week <= d <= week + dt.timedelta(days=6))]
    if outside:
        warnings.append(DataWarning(INFO, 'Configuració', f"Festius fora de la setmana (no afecten): {', '.join(fmt_date(d) for d in outside)}."))


# ---------------------------------------------------------------------------
# Entrada completa
# ---------------------------------------------------------------------------

def load_entrada(path, display_name: str = '') -> Entrada:
    book = Book.open(path)
    book.label = display_name or Path(str(path)).name
    warnings = []
    params, week = _read_config(book, warnings)
    activities, rads = _read_radiologists(book, warnings)
    names = NameIndex(rads)
    blocks = []
    if activities:
        blocks, week = _read_demand(book, activities, names, week, warnings)
    absences = _read_absences(book, names, week, warnings) if week else {}
    if week:
        params.holidays = sorted(params.holidays)
        _cross_checks(rads, activities, blocks, week, params, warnings)
    if book.unresolved:
        cells = ', '.join(f"{sh}!{get_column_letter(c)}{r}" for sh, r, c, _ in book.unresolved[:6])
        warnings.append(DataWarning(AVIS, 'Fitxer', f"Hi ha {len(book.unresolved)} fórmules sense valor calculat ({cells}{'…' if len(book.unresolved) > 6 else ''}): s'han llegit com a buides.",
                                    book.label, "Obre el fitxer amb Excel, desa'l i torna'l a pujar."))
    return Entrada(source=book.label, week=week, params=params, activities=activities, radiologists=rads,
                   blocks=blocks, absences=absences, warnings=warnings)


def describe_radiologist(r: Radiologist, activities: dict) -> str:
    acts = [f"{activities[k].label}{' (P)' if v == 'P' else ''}" for k, v in r.competences.items() if k in activities]
    extra = []
    if r.centres:
        extra.append('centres: ' + ', '.join(r.centres))
    if r.days_off:
        extra.append('no treballa: ' + weekday_list_text(r.days_off))
    return '; '.join([', '.join(acts)] + extra)


__all__ = ['load_entrada', 'parse_competence', 'NameIndex', 'describe_radiologist']
