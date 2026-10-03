"""Regressions de la revisió independent (una prova per troballa)."""
import copy
import datetime as dt
from pathlib import Path

import openpyxl
import pytest

from programada.model import ERROR, AVIS, centre_matches
from programada.normalize import INGR, MIXT, parse_ambit, parse_active, parse_number
from programada.io_input import load_entrada, parse_competence
from programada.templates import build_input
from programada.availability import build_radweeks, parse_absence_code
from programada.planner import plan
from conftest import Factory, WEEK, solve_factory, plan_factory

ROOT = Path(__file__).resolve().parent.parent


def _rad(name, mn=0, mx=50, comps=None, centres='', days_off='', actiu='Sí'):
    return {'Radiòleg': name, 'Actiu aquesta setmana': actiu, 'Mínim setmanal': mn, 'Màxim setmanal': mx,
            'Centres on pot informar': centres, 'Dies sense activitat': days_off, 'competències': comps or {}}


def _dem(centre, wd, act, n, ambit='', fix=''):
    return {'Centre': centre, 'Data': WEEK + dt.timedelta(days=wd) if wd is not None else None, 'Franja': 'Matí',
            'Agenda / equip': 'EQ', 'Activitat': act, 'Exploracions': n, 'Àmbit': ambit, 'Radiòleg fix': fix}


def _book(tmp_path, name='e.xlsx', **kw):
    p = tmp_path / name
    kw.setdefault('week', WEEK)
    build_input(str(p), **kw)
    return p


# 1 ------------------------------------------------------------------------------------
def test_one_exam_block_keeps_the_model_valid():
    f = Factory()
    f.rad('A', 0, 30, {'TC Body': 'S'})
    f.rad('B', 0, 30, {'TC Body': 'S'})
    f.block('TC Body', 1)
    f.block('TC Body', 12, day=1)
    ent, rws, out = solve_factory(f)
    assert out.stats.status == 'OPTIMAL'
    assert not out.unassigned


# 2 ------------------------------------------------------------------------------------
def test_active_column_values():
    for v in ('No', 'No.', 'Baixa', 'No (baixa)', 'Excedència', 'Inactiu'):
        assert parse_active(v) is False, v
    for v in (None, '', 'Sí', 'si', 'x', 'Actiu'):
        assert parse_active(v) is True, v
    assert parse_active('potser') is None


def test_unclear_active_value_is_inactive_with_warning(tmp_path):
    p = _book(tmp_path, radiologists=[_rad('A', comps={'TC Body': 'S'}, actiu='potser'), _rad('B', comps={'TC Body': 'S'})],
              demand=[_dem('H', 0, 'TC Body', 10)])
    ent = load_entrada(str(p))
    assert not ent.radiologists['a'].active
    assert any("'Actiu' = 'potser'" in w.message for w in ent.warnings)


# 3 ------------------------------------------------------------------------------------
def test_mixed_scope_written_in_words():
    for v in ('Ambulatori i ingressat', 'Amb/Ingr', 'CEX + ingressats', 'Mixt'):
        assert parse_ambit(v) == MIXT, v
    assert parse_ambit('Ingressats') == INGR and parse_ambit('urgent') is None


# 4 ------------------------------------------------------------------------------------
def test_centre_matching_respects_numbers_and_letters():
    assert not centre_matches('CAP Badalona 2', 'CAP Badalona 1')
    assert not centre_matches('Edifici A', 'Edifici B')
    assert centre_matches('Bellvitge', 'Hospital Universitari de Bellvitge')
    assert centre_matches("Hospital d'Igualada", 'Hospital Igualada')


def test_all_centres_keyword(tmp_path):
    p = _book(tmp_path, radiologists=[_rad('A', comps={'TC Body': 'S'}, centres='Tots')],
              demand=[_dem('Hospital Demo Sud', 0, 'TC Body', 10)])
    ent = load_entrada(str(p))
    assert ent.radiologists['a'].centres == []


# 5 ------------------------------------------------------------------------------------
def test_fixed_radiologist_absent_that_day_is_substituted():
    f = Factory()
    f.rad('Fix', 0, 50, {'RM Neuro': 'S'})
    f.rad('Altre', 0, 50, {'RM Neuro': 'S'})
    f.absent('Fix', 0)
    f.block('RM Neuro', 16, day=0, fixed='Fix')
    res = plan_factory(f)
    a = next(a for a in res.assignments if a.block_id == 'F001')
    assert a.radiologist_key == 'altre'
    assert any(fl.startswith('Substitueix el radiòleg fix Fix (vacances el 19/10)') for fl in a.flags)
    assert not [w for w in res.validation if w.level == ERROR]


# 6 ------------------------------------------------------------------------------------
def test_title_row_does_not_steal_the_header(tmp_path):
    p = _book(tmp_path, radiologists=[_rad('Alpha, Anna', comps={'TC Body': 'S'})],
              demand=[_dem('H', 0, 'TC Body', 10)], absences={'Alpha, Anna': {0: 'V'}})
    wb = openpyxl.load_workbook(p)
    ws = wb['Absències']
    ws['A1'] = 'Radiòlegs absents setmana 43'
    ws.delete_cols(9)                                   # sense columna Observacions
    wb.save(p)
    ent = load_entrada(str(p))
    assert ent.absences['alphaanna'] == {WEEK: 'V'}


def test_header_below_row_20(tmp_path):
    p = _book(tmp_path, radiologists=[_rad('Alpha, Anna', comps={'TC Body': 'S'})], demand=[_dem('H', 0, 'TC Body', 10)])
    wb = openpyxl.load_workbook(p)
    wb['Radiòlegs'].insert_rows(3, 18)
    wb.save(p)
    ent = load_entrada(str(p))
    assert list(ent.radiologists) == ['alphaanna'] and not ent.has_errors()


# 7, 8, 16 ----------------------------------------------------------------------------
def _app(data: bytes, name='entrada.xlsx'):
    st_testing = pytest.importorskip('streamlit.testing.v1')
    at = st_testing.AppTest.from_file(str(ROOT / 'streamlit_app.py'), default_timeout=180)
    at.session_state['input'] = (name, data)
    at.session_state['input_source'] = 'demo'
    at.run()
    return at


def test_app_accepts_parameters_outside_widget_ranges(tmp_path):
    p = _book(tmp_path, radiologists=[_rad('A', comps={'TC Body': 'S'})], demand=[_dem('H', 0, 'TC Body', 10)],
              params={'Temps màxim de càlcul': 2, 'Màxim de radiòlegs per bloc': 30, 'Mida mínima de lot': 60,
                      'Cost de dividir': 500})
    at = _app(p.read_bytes())
    assert not at.exception


def test_rerun_after_assigning_does_not_claim_stale_parameters():
    from programada.demo import build_demo_input
    import tempfile
    import os
    with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as fh:
        path = fh.name
    build_demo_input(path)
    data = Path(path).read_bytes()
    os.unlink(path)
    at = _app(data)
    at.button[0].click().run()                          # Assignar
    at.multiselect[0].select('Alpha, Anna').run()       # qualsevol interacció posterior
    assert not at.exception
    assert not any('Has canviat els paràmetres' in i.value for i in at.info)


def test_corrupt_upload_shows_an_error():
    at = _app(b'aixo no es un excel', 'malmes.xlsx')
    assert not at.exception
    assert any('No es pot llegir' in e.value for e in at.error)


# 9 ------------------------------------------------------------------------------------
def test_activity_columns_are_not_decided_by_annotations(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Radiòlegs'
    ws.append(['Radiòleg', 'Màxim setmanal', 'Dedicació', 'RM Cardio', 'TC Body'])
    ws.append(['A', 40, 1, 'S (només matins)', '✓'])
    ws.append(['B', 40, '0,8', 'S*', 'S'])
    ws.append(['C', 40, '0,5', 'En formació', None])
    ws = wb.create_sheet('Demanda')
    ws.append(['Centre', 'Data', 'Activitat', 'Exploracions'])
    ws.append(['H', WEEK, 'RM Cardio', 10])
    p = tmp_path / 'act.xlsx'
    wb.save(p)
    ent = load_entrada(str(p))
    assert set(ent.activities) == {'rmcardio', 'tcbody'}           # 'Dedicació' no és activitat
    assert ent.radiologists['a'].competences == {'rmcardio': 'S', 'tcbody': 'S'}
    assert ent.radiologists['b'].competences == {'rmcardio': 'S', 'tcbody': 'S'}
    assert 'rmcardio' not in ent.radiologists['c'].competences
    assert any("'En formació'" in w.message for w in ent.warnings)
    assert parse_competence('P - preferent') == 'P' and parse_competence('✔') == 'S'


# 10 -----------------------------------------------------------------------------------
def test_high_split_cost_never_leaves_coverable_demand_pending():
    f = Factory(split_cost_exams=100)
    f.rad('A', 0, 10, {'TC Body': 'S'})
    f.rad('B', 0, 10, {'TC Body': 'S'})
    f.block('TC Body', 12)
    f.block('TC Body', 2, day=1)
    ent, rws, out = solve_factory(f)
    assert not out.unassigned


# 11, 12 -------------------------------------------------------------------------------
def test_pending_reasons_use_plan_parameters_and_real_margins():
    f = Factory(min_chunk=5)
    f.rad('A', 0, 4, {'TC Body': 'S'})
    f.rad('B', 0, 4, {'TC Body': 'S'})
    f.block('TC Body', 8)
    res = plan_factory(f)
    txt = ' '.join(res.pendings[0].reasons)
    assert 'poc marge' in txt and 'Capacitat esgotada' not in txt
    # Paràmetres del càlcul, no els del fitxer
    f2 = Factory()
    for n in ('A', 'B', 'C'):
        f2.rad(n, 0, 10, {'TC Body': 'S'})
    f2.block('TC Body', 30)
    p = copy.deepcopy(f2.params)
    p.max_rads_per_block = 2                            # la Configuració del fitxer en diu 4
    res2 = plan(f2.entrada(), p)
    assert any('màxim de radiòlegs (2)' in r for r in res2.pendings[0].reasons)


def test_fixed_without_competence_is_not_offered_as_candidate():
    f = Factory()
    f.rad('Fix', 0, 50, {'RM Neuro': 'S'})
    f.rad('Un', 0, 10, {'TC Body': 'S'})
    f.block('TC Body', 30, fixed=['Fix', 'Un'])
    res = plan_factory(f)
    assert all('Fix' not in c for p in res.pendings for c in p.candidates)


# 13 -----------------------------------------------------------------------------------
def test_undated_inpatient_volume_has_priority():
    f = Factory()
    f.rad('A', 0, 10, {'TC Body': 'P', 'RM Body': 'S'})
    f.block('TC Body', 10, day=0)
    f.block('RM Body', 10, day=None, scope=INGR)
    ent, rws, out = solve_factory(f)
    assert out.x.get(('F002', 'a')) == 10


# 14 -----------------------------------------------------------------------------------
def test_days_off_covering_the_whole_week_are_ignored():
    f = Factory()
    f.rad('A', 0, 30, {'TC Body': 'S'}, days_off='Dl-Dv')
    rws, warns = build_radweeks(f.entrada())
    assert rws['a'].cap == 30
    assert any('no es tenen en compte' in w.message for w in warns)


# 15 -----------------------------------------------------------------------------------
def test_formulas_without_cached_value_are_reported(tmp_path):
    p = _book(tmp_path, radiologists=[_rad('A', comps={'TC Body': 'S'})], demand=[_dem('H', 0, 'TC Body', 10)])
    wb = openpyxl.load_workbook(p)
    wb['Demanda']['B5'] = "='Configuració'!$B$5+4"
    wb.save(p)
    ent = load_entrada(str(p))
    assert any('fórmules sense valor calculat' in w.message for w in ent.warnings)


# 17 -----------------------------------------------------------------------------------
def test_duplicate_radiologist_is_a_warning_not_a_block(tmp_path):
    p = _book(tmp_path, radiologists=[_rad('A', comps={'TC Body': 'S'}), _rad('A', 0, 10, comps={'TC Body': 'S'})],
              demand=[_dem('H', 0, 'TC Body', 10)])
    ent = load_entrada(str(p))
    assert not ent.has_errors()
    assert any('apareix dues vegades' in w.message and w.level == AVIS for w in ent.warnings)


# Altres casos de la revisió --------------------------------------------------------------
def test_numbers_dates_and_codes_as_people_write_them(tmp_path):
    assert parse_number('1.000') == 1000 and parse_number('1,234.5') == 1234.5 and parse_number('12,5') == 12.5
    assert parse_absence_code('guàrdia ja assignada') == ('X', True)
    assert parse_absence_code('Guàrdia') == ('G', False)
    p = _book(tmp_path, radiologists=[_rad('A', comps={'TC Body': 'S'})],
              demand=[_dem('H', 0, 'TC Body', '2,5'), _dem('H', 0, 'TC Body', '0,4')])
    wb = openpyxl.load_workbook(p)
    wb['Demanda']['B5'] = 'Dilluns i dimarts'
    wb.save(p)
    ent = load_entrada(str(p))
    assert [b.n_exams for b in ent.blocks] == [3]
    assert ent.blocks[0].date is None
    assert any('diversos dies' in w.message for w in ent.warnings)


# Segona revisió ------------------------------------------------------------------------
def test_absent_fixed_radiologist_still_covers_when_nobody_else_can():
    f = Factory()
    f.rad('Fix, Un', 0, 50, {'TC Body': 'S'})
    f.absent('Fix, Un', 0)
    f.block('TC Body', 16, day=0, fixed='Fix, Un')
    res = plan_factory(f)
    assert not res.pendings
    a = res.assignments[0]
    assert a.radiologist_key == 'fixun' and "l'informarà a partir del Dt 20/10" in ' '.join(a.flags)
    assert not [w for w in res.validation if w.level == ERROR]


def test_weekend_and_holiday_agendas_keep_their_fixed_radiologist():
    f = Factory()
    f.params.holidays = [WEEK + dt.timedelta(days=2)]
    f.rad('Fix', 0, 50, {'RM Neuro': 'S'})
    f.rad('Altre', 0, 50, {'RM Neuro': 'S'})
    f.block('RM Neuro', 12, day=5, fixed='Fix')        # dissabte
    f.block('RM Neuro', 12, day=2, fixed='Fix')        # festiu
    res = plan_factory(f)
    assert {a.block_id: a.radiologist_key for a in res.assignments} == {'F001': 'fix', 'F002': 'fix'}
    assert not any('Substitueix' in fl for a in res.assignments for fl in a.flags)


def test_abbreviations_and_initials_still_match(tmp_path):
    assert centre_matches('H. Bellvitge', 'Hospital Universitari de Bellvitge')
    assert centre_matches('H.U. Bellvitge', 'Hospital Universitari de Bellvitge')
    p = _book(tmp_path, radiologists=[_rad('Alpha, Anna', comps={'TC Body': 'S'}), _rad('Beta, Bernat', comps={'TC Body': 'S'})],
              demand=[_dem('H', 1, 'TC Body', 10, ambit='Ingressat', fix='A. Alpha')], absences={'Beta, B.': {1: 'V'}})
    ent = load_entrada(str(p))
    assert ent.blocks[0].fixed_keys == ['alphaanna']
    assert ent.absences == {'betabernat': {WEEK + dt.timedelta(days=1): 'V'}}


def test_active_values_are_whole_words():
    for v in ('Només matins', 'Nou', 'Normal', 'Incorporació', 'Interí'):
        assert parse_active(v) is None, v


def test_negated_inpatients_are_ambulatory():
    assert parse_ambit('Sense ingressats') == 'Ambulatori'
    assert parse_ambit('No ingressat') == 'Ambulatori'


def test_dash_means_does_not_do_it(tmp_path):
    assert parse_competence('-') == '' and parse_competence('—') == ''
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Radiòlegs'
    ws.append(['Radiòleg', 'Màxim setmanal', 'Neuro', 'Body'])
    ws.append(['A', 40, 'S', '-'])
    ws.append(['B', 40, '-', 'P'])
    ws = wb.create_sheet('Demanda')
    ws.append(['Centre', 'Data', 'Activitat', 'Exploracions'])
    ws.append(['H', WEEK, 'Neuro', 10])
    p = tmp_path / 'guions.xlsx'
    wb.save(p)
    ent = load_entrada(str(p))
    assert set(ent.activities) == {'neuro', 'body'} and not ent.has_errors()
    assert ent.radiologists['b'].competences == {'body': 'P'}


def test_numeric_column_after_the_template_activities_is_ignored(tmp_path):
    p = _book(tmp_path, radiologists=[_rad('A', comps={'TC Body': 'S'})], demand=[_dem('H', 0, 'TC Body', 10)])
    wb = openpyxl.load_workbook(p)
    ws = wb['Radiòlegs']
    col = ws.max_column + 1
    ws.cell(row=4, column=col, value='Dedicació')
    ws.cell(row=5, column=col, value=0.8)
    wb.save(p)
    ent = load_entrada(str(p))
    assert 'dedicacio' not in ent.activities


def test_several_weekdays_in_date_text(tmp_path):
    p = _book(tmp_path, radiologists=[_rad('A', comps={'TC Body': 'S'})], demand=[_dem('H', 0, 'TC Body', 10)])
    wb = openpyxl.load_workbook(p)
    wb['Demanda']['B5'] = 'Dilluns i Dimarts (tarda)'
    wb.save(p)
    ent = load_entrada(str(p))
    assert ent.blocks[0].date is None and any('diversos dies' in w.message for w in ent.warnings)


# Tercera revisió ------------------------------------------------------------------------
def test_bare_letters_are_part_of_the_centre_name():
    assert not centre_matches('CAP Terrassa A', 'CAP Terrassa Nord')
    assert not centre_matches('Edifici A', 'Edifici Principal')
    assert centre_matches('H.U.Bellvitge', 'Hospital Universitari de Bellvitge')


def test_suspended_pact_prefers_who_is_there_even_against_obligation():
    f = Factory()
    f.rad('Fix', 40, 60, {'RM Neuro': 'S'})
    f.rad('Present', 0, 30, {'RM Neuro': 'S'})
    f.absent('Fix', 0)
    f.block('RM Neuro', 16, day=0, fixed='Fix')
    res = plan_factory(f)
    a = next(a for a in res.assignments if a.block_id == 'F001')
    assert a.radiologist_key == 'present'
    assert any('Substitueix el radiòleg fix Fix' in fl for fl in a.flags)


def test_cap_primary_care_is_not_a_negation():
    assert parse_ambit('CAP/Hosp') in (INGR, MIXT)
    assert parse_ambit('CAP + hospitalitzats') in (INGR, MIXT)
    assert parse_ambit('Cap ingressat') == 'Ambulatori'


def test_numbers_are_not_evidence_of_an_activity_column(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Radiòlegs'
    ws.append(['Radiòleg', 'Màxim setmanal', 'TC Body', 'Dedicació'])
    for i, ded in enumerate([1, 1, 1, '0,8']):
        ws.append([f'R{i}', 40, 'S', ded])
    ws = wb.create_sheet('Demanda')
    ws.append(['Centre', 'Data', 'Activitat', 'Exploracions'])
    ws.append(['H', WEEK, 'TC Body', 10])
    p = tmp_path / 'ded.xlsx'
    wb.save(p)
    ent = load_entrada(str(p))
    assert set(ent.activities) == {'tcbody'}
