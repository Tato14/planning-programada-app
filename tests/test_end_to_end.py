"""Flux complet amb la demo: Excel d'entrada -> assignació -> Excel de sortida, CLI i app."""
from pathlib import Path

import openpyxl
import pytest

from programada.demo import build_demo_input
from programada.io_input import load_entrada
from programada.planner import plan
from programada.writer import write_plan
from programada.model import ERROR
from programada import cli

ROOT = Path(__file__).resolve().parent.parent


@pytest.fixture(scope='module')
def demo(tmp_path_factory):
    d = tmp_path_factory.mktemp('demo')
    p = d / 'entrada.xlsx'
    build_demo_input(str(p))
    ent = load_entrada(str(p))
    return d, p, ent, plan(ent)


def test_demo_plan_is_valid_and_explained(demo):
    d, p, ent, res = demo
    assert res.stats.status in ('OPTIMAL', 'FEASIBLE')
    assert not [w for w in res.validation if w.level == ERROR], [w.message for w in res.validation]
    demand = sum(b.n_exams for b in res.blocks)
    assert sum(a.n_exams for a in res.assignments) + sum(x.n_exams for x in res.pendings) == demand
    bmap = res.blocks_by_id()
    # Activitat fora del catàleg i cardio amb més demanda que capacitat: pendents amb motiu
    reasons = {bmap[x.block_id].activity_label: ' '.join(x.reasons) for x in res.pendings}
    assert 'no és cap de les columnes' in reasons['RM Mama']
    assert 'Capacitat esgotada' in reasons['RM Cardio']
    # La radiòloga fixa de vacances: substituïda i marcat
    sub = [a for a in res.assignments if any('Substitueix el radiòleg fix Nu, Núria' in f for f in a.flags)]
    assert sub and all(a.radiologist_key != 'nunuria' for a in res.assignments)
    # El fix escrit 'Anna Alpha' es respecta
    fixed_alpha = [a for a in res.assignments if a.radiologist_key == 'alphaanna' and a.fixed]
    assert fixed_alpha and fixed_alpha[0].kind.startswith('Fixa')
    # Ingressats: només a qui treballa aquell dia
    for a in res.assignments:
        b = bmap[a.block_id]
        if b.scope != 'Ambulatori' and b.date:
            assert res.radweeks[a.radiologist_key].available_on(b.date)
    # Qui només cobreix guàrdies (Màxim 0) i l'inactiu no reben res
    assert all(a.radiologist_key not in ('thetahector', 'omicronpau') for a in res.assignments)


def test_output_workbook_has_named_tables(demo, tmp_path):
    d, p, ent, res = demo
    out = tmp_path / 'pla.xlsx'
    write_plan(res, str(out))
    wb = openpyxl.load_workbook(out)
    tables = {t for ws in wb.worksheets for t in ws.tables}
    assert {'tAssignacions', 'tRadiolegs', 'tPendents', 'tAvisos', 'tParametres', 'tResumActivitat', 'tResumCentre'} <= tables
    ws = wb['Assignacions']
    hdr = [c.value for c in ws[1]]
    n_col = hdr.index('Exploracions') + 1
    assert sum(ws.cell(row=r, column=n_col).value or 0 for r in range(2, ws.max_row + 1)) == sum(a.n_exams for a in res.assignments)


def test_cli(demo, tmp_path, capsys):
    d, p, ent, res = demo
    out = tmp_path / 'cli.xlsx'
    assert cli.main(['assignar', str(p), '-o', str(out), '--temps', '10']) == 0
    assert out.exists() and 'exploracions assignades' in capsys.readouterr().out
    assert cli.main(['plantilla', '-o', str(tmp_path / 'p.xlsx')]) == 0
    bad = tmp_path / 'buida.xlsx'
    cli.main(['plantilla', '-o', str(bad)])
    assert cli.main(['assignar', str(bad)]) == 1


def test_app_runs_with_demo():
    st_testing = pytest.importorskip('streamlit.testing.v1')
    at = st_testing.AppTest.from_file(str(ROOT / 'streamlit_app.py'), default_timeout=180)
    at.run()
    assert not at.exception
    at.sidebar.button[0].click().run()          # Carregar la demo
    assert not at.exception
    at.button[0].click().run()                  # Assignar
    assert not at.exception
    assert any('Validació independent superada' in s.value for s in at.success)
