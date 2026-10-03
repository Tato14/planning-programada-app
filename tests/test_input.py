"""Lectura del fitxer d'entrada: plantilla, demo i variants que escriuen les persones."""
import datetime as dt

import openpyxl

from programada.io_input import load_entrada, parse_competence
from programada.templates import build_input
from programada.demo import build_demo_input, DEMO_WEEK
from programada.model import ERROR, AVIS
from programada.normalize import INGR, MIXT, AMB

WEEK = DEMO_WEEK


def _msgs(ent, level=None):
    return [w.message for w in ent.warnings if level is None or w.level == level]


def test_blank_template_reads_without_crashing(tmp_path):
    p = tmp_path / 'buida.xlsx'
    build_input(str(p))
    ent = load_entrada(str(p))
    assert len(ent.activities) == 10
    assert not ent.radiologists and not ent.blocks
    assert ent.has_errors()          # no hi ha radiòlegs ni setmana: no es pot assignar


def test_demo_input(tmp_path):
    p = tmp_path / 'demo.xlsx'
    build_demo_input(str(p))
    ent = load_entrada(str(p))
    assert not ent.has_errors(), _msgs(ent, ERROR)
    assert ent.week == WEEK
    assert len(ent.blocks) == 28 and sum(b.n_exams for b in ent.blocks) == 447
    assert any("'RM Mama'" in m for m in _msgs(ent, AVIS))
    # Noms d'Absències per fórmula des de Radiòlegs (fitxer sense desar amb Excel)
    assert len(ent.absences['nunuria']) == 5 and ent.absences['gammacarla'][WEEK + dt.timedelta(days=4)] == 'M'
    # 'Anna Alpha' és Alpha, Anna
    b = next(b for b in ent.blocks if b.fixed_names)
    assert b.fixed_keys == ['alphaanna']
    assert {b.scope for b in ent.blocks} == {AMB, INGR, MIXT}
    assert any(b.date is None for b in ent.blocks)          # volum de la setmana
    assert ent.radiologists['epsilonelena'].days_off == [3, 4]
    assert ent.radiologists['alphaanna'].centres == ['Hospital Demo Nord', 'Hospital Demo Sud']
    assert not ent.radiologists['omicronpau'].active


def _variant_book(path, with_config=True):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Configuracio'
    if with_config:
        ws.append(['Paràmetre', 'Valor'])
        ws.append(['Setmana', '21/10/2026'])                 # dimecres -> dilluns 19/10
        ws.append(['Festius', '23/10/2026, 01/01/2027'])
        ws.append(['Ocupació objectiu', '0,7'])
        ws.append(['Lot mínim', 3])
    ws = wb.create_sheet('Radiolegs')
    ws.append(['Llista de radiòlegs'])
    ws.append([])
    ws.append(['Nom', 'Capacitat', 'Obligació', 'Vinculació', 'rm-neuro', 'TC body', 'Centres', 'Dies que no treballa'])
    ws.append(['Alpha, Anna', '40', '20', 'TD', 'Sí', 'pref', 'Nord; Sud', 'divendres'])
    ws.append(['Beta, Bernat', 30, None, 'TD', 'x', None, None, None])
    ws.append(['Beta, Berta', 30, None, 'Extern', None, 'S', None, 'Dl i Dt'])
    ws.append(['Alpha, Anna', 10, 0, 'TD', 'S', None, None, None])
    ws = wb.create_sheet('Demanda setmanal')
    ws.append(['Hospital', 'Dia', 'Torn', 'Actividad', 'Nº exploracions', 'Tipus de pacient', 'Radiòleg assignat'])
    ws.append(['Hospital Nord', 'Dimecres', '15-21', 'RM neuro', 12, 'hosp', 'Alpha'])
    ws.append(['Hospital Nord', dt.date(2026, 10, 20), 'M', 'body TC', '7,6', None, 'Beta'])
    ws.append(['Hospital Sud', dt.date(2026, 10, 30), 'Matí', 'TC Body', 5, None, None])
    ws.append(['Hospital Sud', None, None, 'TC Body', 0, None, None])
    ws.append(['Hospital Sud', None, None, 'Eco', 4, 'urgent', 'Gamma'])
    ws = wb.create_sheet('Vacances')
    ws.append(['Radiòleg', 'Dl', 'Dt', 'Dc', 'Dj', 'Dv'])
    ws.append(['Anna Alpha', 'vacances', None, None, 'Z', None])
    ws.append(['Algú altre', 'V', None, None, None, None])
    wb.save(path)


def test_tolerant_reading_of_variants(tmp_path):
    p = tmp_path / 'variants.xlsx'
    _variant_book(str(p))
    ent = load_entrada(str(p))
    m = _msgs(ent)
    assert ent.week == WEEK and ent.params.target_utilisation == 0.7 and ent.params.min_chunk == 3
    assert ent.params.holidays == [dt.date(2026, 10, 23), dt.date(2027, 1, 1)]
    assert set(ent.activities) == {'rmneuro', 'tcbody'}                       # 'Vinculació' no és activitat
    assert any("Columna 'Vinculació' ignorada" in x for x in m)
    a = ent.radiologists['alphaanna']
    assert a.weekly_max == 40 and a.weekly_min == 20 and a.competences == {'rmneuro': 'S', 'tcbody': 'P'}
    assert a.centres == ['Nord', 'Sud'] and a.days_off == [4]
    assert any('apareix dues vegades' in x for x in m)
    assert ent.radiologists['betaberta'].days_off == [0, 1]                  # 'Dl i Dt'
    b1, b2 = ent.blocks[0], ent.blocks[1]
    assert b1.date == dt.date(2026, 10, 21) and b1.franja == 'Tarda' and b1.scope == INGR and b1.activity_key == 'rmneuro'
    assert b1.fixed_keys == ['alphaanna']                                     # 'Alpha' només és Alpha, Anna
    assert b2.activity_key == 'tcbody' and b2.n_exams == 8 and b2.franja == 'Matí'
    assert b2.fixed_keys == []                                                # 'Beta' és ambigu
    assert any("'Beta' pot ser" in x for x in m)
    assert any('fora de la setmana' in x for x in m)
    assert any("nombre d'exploracions vàlid" in x for x in m)
    eco = next(b for b in ent.blocks if b.activity_label == 'Eco')
    assert eco.activity_key == '' and eco.scope == AMB
    assert any("Àmbit 'urgent'" in x for x in m) and any("'Gamma' no és a la pestanya" in x for x in m)
    # Absències: paraula llarga, codi desconegut (avís) i nom que no hi és
    assert ent.absences['alphaanna'] == {WEEK: 'V', WEEK + dt.timedelta(days=3): 'A'}
    assert any("codi 'Z' no reconegut" in x for x in m)
    assert any("'Algú altre' no és a la pestanya" in x for x in m)


def test_week_is_inferred_without_configuration(tmp_path):
    p = tmp_path / 'sense_config.xlsx'
    _variant_book(str(p), with_config=False)
    ent = load_entrada(str(p))
    assert ent.week == WEEK
    assert any('es dedueix de les dates' in x for x in _msgs(ent, AVIS))


def test_competence_values():
    assert parse_competence('Sí') == 'S' and parse_competence('X') == 'S' and parse_competence('P') == 'P'
    assert parse_competence('preferent') == 'P' and parse_competence('') == '' and parse_competence('No') == ''
    assert parse_competence('TD') is None


def test_centre_token_matching_several_centres_warns(tmp_path):
    p = tmp_path / 'centres.xlsx'
    build_input(str(p), week=WEEK,
                radiologists=[{'Radiòleg': 'Alpha, Anna', 'Màxim setmanal': 30, 'Centres on pot informar': 'Hospital',
                               'competències': {'TC Body': 'S'}}],
                demand=[{'Centre': 'Hospital Demo Nord', 'Data': WEEK, 'Activitat': 'TC Body', 'Exploracions': 5},
                        {'Centre': 'Hospital Demo Sud', 'Data': WEEK, 'Activitat': 'TC Body', 'Exploracions': 5}])
    ent = load_entrada(str(p))
    assert any("coincideix amb 2 centres" in x for x in _msgs(ent, AVIS))
    assert not ent.has_errors()
